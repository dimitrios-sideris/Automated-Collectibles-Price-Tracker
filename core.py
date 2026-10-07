"""Core inventory, price-history and database logic for the public tracker."""
from __future__ import annotations

# Standard-library imports: no pip installation is needed for these.
import json              # parse Cardmarket JSON files
import os                # build/check/create file and directory paths
import sqlite3           # local database for matches and price history
import ssl               # HTTPS certificate context for urllib
import tempfile          # safe temporary file for the large daily price guide
import urllib.parse      # inspect and validate download URLs
import urllib.request    # HTTP/HTTPS downloads without an external requests package
from datetime import datetime
from typing import Iterable

# The project's only third-party package. It is used to read the Excel workbook.
from openpyxl import load_workbook

# -----------------------------------------------------------------------------
# Project paths and input files
# -----------------------------------------------------------------------------
# All file locations are built with os.path relative to this source file.
# That means the project can be moved or cloned to another folder without
# changing hard-coded absolute paths.
ROOT = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(ROOT, "data")
DB_PATH = os.path.join(DATA_DIR, "yugioh.sqlite")
PRIVATE_INVENTORY = os.path.join(DATA_DIR, "inventory.xlsx")
EXAMPLE_INVENTORY = os.path.join(DATA_DIR, "inventory.example.xlsx")
SHEET_NAME = "Triple_Loose"

# -----------------------------------------------------------------------------
# Cardmarket daily price-guide download
# -----------------------------------------------------------------------------
# Game ID 3 is Yu-Gi-Oh on Cardmarket.  The size limits below are defensive:
# they help reject an obviously wrong/truncated response before parsing it.
CARDMARKET_HOST = "downloads.s3.cardmarket.com"
PRICE_GUIDE_URL = (
    "https://downloads.s3.cardmarket.com/"
    "productCatalog/priceGuide/price_guide_3.json"
)
MAX_DOWNLOAD_BYTES = 80 * 1024 * 1024
MIN_DOWNLOAD_BYTES = 1 * 1024 * 1024

# Columns the pricing pipeline truly needs.  The other Excel columns (Printed
# Name, Location, Notes, etc.) are useful metadata but are intentionally optional.
REQUIRED_COLUMNS = {
    "Card ID",
    "English Name",
    "Card Code",
    "Rarity",
    "Quantity",
}


# Inventory selection is centralized here.  A local private workbook, when
# present, takes priority; otherwise the public example workbook is used.
def inventory_path() -> str:
    """Prefer an optional private inventory; otherwise use the public workbook."""
    if os.path.exists(PRIVATE_INVENTORY):
        return PRIVATE_INVENTORY
    if os.path.exists(EXAMPLE_INVENTORY):
        return EXAMPLE_INVENTORY
    raise FileNotFoundError(
        "No inventory workbook found. Expected data/inventory.xlsx or "
        "data/inventory.example.xlsx."
    )


# Small Excel-cleaning helper used while reading cells.  Empty strings become
# None so later code does not need to distinguish several kinds of "missing".
def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


# Read the human-maintained Excel inventory and convert each row into a normal
# Python dictionary.  No Cardmarket Product ID is expected in the workbook;
# matching is deliberately handled later by matcher.py.
def load_inventory(path: str | None = None) -> list[dict]:
    """Read Triple_Loose.  Users never need to provide a marketplace ID."""
    # Resolve the selected workbook to an absolute string path for clearer errors/logs.
    path = os.path.abspath(path or inventory_path())

    # read_only=True keeps memory use low; data_only=True reads displayed values
    # instead of Excel formulas when a workbook happens to contain formulas.
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        # Fail early if the expected inventory sheet is not present.
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Workbook has no sheet named {SHEET_NAME!r}: {path}")
        ws = wb[SHEET_NAME]

        # Search for the header instead of assuming it is always Excel row 1.
        # This tolerates a title/blank row above the table.
        header_row = None
        header = None
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [_clean(v) for v in row]
            if "Card ID" in values and "Card Code" in values:
                header_row = row_index
                header = values
                break
        if header_row is None or header is None:
            raise ValueError("Could not locate the inventory header row.")

        # Validate only the columns needed by the program; metadata columns stay optional.
        missing = REQUIRED_COLUMNS - {v for v in header if v}
        if missing:
            raise ValueError(
                "Inventory is missing required columns: " + ", ".join(sorted(missing))
            )

        # Convert column names to numeric positions once, then reuse that map for every row.
        index = {name: i for i, name in enumerate(header) if name}
        rows: list[dict] = []
        seen_ids: set[str] = set()

        # Read every inventory row below the header and normalize it into a dict.
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            values = [_clean(v) for v in values]

            def get(name, default=None):
                pos = index.get(name)
                return values[pos] if pos is not None and pos < len(values) else default

            # Card ID is the stable primary key used throughout SQLite and the GUI.
            card_id = get("Card ID")
            if not card_id:
                continue
            card_id = str(card_id)
            if card_id in seen_ids:
                raise ValueError(f"Duplicate Card ID: {card_id}")
            seen_ids.add(card_id)

            # Quantity must be an integer because portfolio value = quantity × unit price.
            quantity = get("Quantity", 1)
            try:
                quantity = int(quantity or 0)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Invalid Quantity for {card_id}: {quantity!r}") from exc
            if quantity < 0:
                raise ValueError(f"Quantity cannot be negative for {card_id}.")

            rows.append(
                {
                    "card_id": card_id,
                    "printed_name": get("Printed Name"),
                    "english_name": get("English Name"),
                    "language": get("Language"),
                    "card_code": get("Card Code"),
                    "rarity": get("Rarity"),
                    "edition": get("Edition"),
                    "quantity": quantity,
                    "condition": get("Condition"),
                    "location": get("Location"),
                    "notes": get("Notes"),
                }
            )

        if not rows:
            raise ValueError("Triple_Loose contains no cards.")
        return rows
    finally:
        # Always close the Excel file, including when validation raises an exception.
        wb.close()


# -----------------------------------------------------------------------------
# SQLite database setup
# -----------------------------------------------------------------------------
# SQLite is the project's local memory: it stores matching decisions, manual
# overrides, daily product prices and daily total portfolio values.
def connect(db_path: str | None = None) -> sqlite3.Connection:
    # Make sure the data directory exists before SQLite tries to create/open the file.
    os.makedirs(DATA_DIR, exist_ok=True)

    # sqlite3 accepts a normal path string, so pathlib is not needed anywhere here.
    con = sqlite3.connect(db_path or DB_PATH)

    # sqlite3.Row lets query results be accessed by column name: row["card_id"].
    con.row_factory = sqlite3.Row

    # Enable SQLite foreign-key checking for any present/future relational constraints.
    con.execute("PRAGMA foreign_keys = ON")
    return con


# Helper used by the lightweight database migration in init_db().
def _columns(con: sqlite3.Connection, table: str) -> set[str]:
    return {row[1] for row in con.execute(f"PRAGMA table_info({table})")}


# Create all core tables if they do not exist.  This function is safe to call
# repeatedly, which is why both the updater and GUI call it at startup.
def init_db(con: sqlite3.Connection) -> None:
    # executescript() is convenient for creating several tables in one SQL block.
    # CREATE TABLE IF NOT EXISTS makes this safe on every program start.
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS inventory (
            card_id TEXT PRIMARY KEY,
            printed_name TEXT,
            english_name TEXT NOT NULL,
            language TEXT,
            card_code TEXT,
            rarity TEXT,
            edition TEXT,
            quantity INTEGER NOT NULL,
            condition TEXT,
            location TEXT,
            notes TEXT,
            product_id INTEGER,
            match_status TEXT,
            match_reason TEXT,
            match_source TEXT,
            match_confidence REAL,
            ygopro_set_name TEXT,
            cardmarket_expansion_id INTEGER,
            candidate_products TEXT
        );

        CREATE TABLE IF NOT EXISTS price_history (
            product_id INTEGER NOT NULL,
            snapshot_date TEXT NOT NULL,
            created_at TEXT NOT NULL,
            avg REAL, low REAL, trend REAL, avg1 REAL, avg7 REAL, avg30 REAL,
            avg_foil REAL, low_foil REAL, trend_foil REAL,
            avg1_foil REAL, avg7_foil REAL, avg30_foil REAL,
            benchmark REAL,
            benchmark_source TEXT,
            PRIMARY KEY (product_id, snapshot_date)
        );

        CREATE TABLE IF NOT EXISTS portfolio_history (
            snapshot_date TEXT PRIMARY KEY,
            created_at TEXT NOT NULL,
            total_value REAL NOT NULL,
            priced_rows INTEGER NOT NULL,
            unpriced_rows INTEGER NOT NULL,
            total_units INTEGER NOT NULL
        );
        """
    )

    # Lightweight migration for anyone who ran the earlier minimal edition.
    required = {
        "match_status": "TEXT",
        "match_reason": "TEXT",
        "match_source": "TEXT",
        "match_confidence": "REAL",
        "ygopro_set_name": "TEXT",
        "cardmarket_expansion_id": "INTEGER",
        "candidate_products": "TEXT",
    }
    # Older local databases may miss newer matching columns. Add only what is absent.
    existing = _columns(con, "inventory")
    for name, sql_type in required.items():
        if name not in existing:
            con.execute(f"ALTER TABLE inventory ADD COLUMN {name} {sql_type}")
    con.commit()

    # Import here (instead of at module import time) to avoid a circular import.
    from matcher import init_match_tables

    # matcher.py owns its own cache/manual-override tables.
    init_match_tables(con)


# Replace the current inventory table with the latest Excel rows *after* matching
# metadata has been attached.  Historical price tables are not deleted here.
def sync_inventory(con: sqlite3.Connection, rows: Iterable[dict]) -> None:
    # Materialize the iterable once because executemany() will consume it.
    rows = list(rows)

    # `with con:` creates a transaction: either the whole inventory mirror is replaced
    # successfully or SQLite rolls it back on failure.
    with con:
        con.execute("DELETE FROM inventory")
        con.executemany(
            """
            INSERT INTO inventory (
                card_id, printed_name, english_name, language, card_code, rarity,
                edition, quantity, condition, location, notes, product_id,
                match_status, match_reason, match_source, match_confidence,
                ygopro_set_name, cardmarket_expansion_id, candidate_products
            ) VALUES (
                :card_id, :printed_name, :english_name, :language, :card_code, :rarity,
                :edition, :quantity, :condition, :location, :notes, :product_id,
                :match_status, :match_reason, :match_source, :match_confidence,
                :ygopro_set_name, :cardmarket_expansion_id, :candidate_products
            )
            """,
            rows,
        )


# -----------------------------------------------------------------------------
# Secure download of Cardmarket's current price guide
# -----------------------------------------------------------------------------
# Redirects are validated as well, so the downloader cannot silently follow an
# unexpected host/path if the remote server behavior changes.
def _validate_price_url(url: str) -> None:
    # Parse the URL into scheme/hostname/path and allow only the expected HTTPS source.
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme.lower() != "https" or parsed.hostname != CARDMARKET_HOST:
        raise ValueError(f"Refusing unexpected download URL: {url}")
    if not parsed.path.endswith("/price_guide_3.json"):
        raise ValueError(f"Unexpected Cardmarket price-guide path: {parsed.path}")


class _SafeRedirect(urllib.request.HTTPRedirectHandler):
    # urllib follows redirects automatically. Validate the redirect target first so a
    # changed remote endpoint cannot silently send the downloader somewhere unexpected.
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        _validate_price_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download_price_guide() -> str:
    """Download Cardmarket's public Yu-Gi-Oh price guide to a temporary file."""
    os.makedirs(DATA_DIR, exist_ok=True)
    _validate_price_url(PRICE_GUIDE_URL)

    # Build an HTTP request with a descriptive User-Agent and ask for JSON.
    request = urllib.request.Request(
        PRICE_GUIDE_URL,
        headers={"User-Agent": "YuGiOhSinglesPriceTracker/2.0", "Accept": "application/json"},
    )
    # Use Python's default trusted certificate store and our redirect validator.
    opener = urllib.request.build_opener(
        _SafeRedirect(), urllib.request.HTTPSHandler(context=ssl.create_default_context())
    )

    # Download to a temporary file because the full price guide is much larger than
    # the subset of products we eventually keep in SQLite.
    handle = tempfile.NamedTemporaryFile(
        mode="wb", suffix=".json", prefix="cardmarket_", dir=DATA_DIR, delete=False
    )
    path = handle.name
    try:
        total = 0
        with handle:
            with opener.open(request, timeout=90) as response:
                # Validate the final URL too, after any redirects.
                _validate_price_url(response.geturl())
                # Stream in 1 MiB chunks rather than loading the entire file into RAM.
                while True:
                    chunk = response.read(1024 * 1024)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > MAX_DOWNLOAD_BYTES:
                        raise ValueError("Cardmarket price-guide download exceeded safety limit.")
                    handle.write(chunk)
        # A very small file is probably an error page or incomplete response.
        if os.path.getsize(path) < MIN_DOWNLOAD_BYTES:
            raise ValueError("Downloaded Cardmarket price guide is suspiciously small.")
        return path
    except Exception:
        # Remove a partial/bad temporary file if downloading or validation failed.
        if os.path.exists(path):
            os.remove(path)
        raise


# -----------------------------------------------------------------------------
# Price selection and price-guide parsing
# -----------------------------------------------------------------------------
def _number(value):
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# Choose one simple unit-value metric for the GUI/portfolio.  We prefer trend,
# then progressively fall back to Cardmarket averages and finally low.
def _benchmark(row: dict) -> tuple[float | None, str | None]:
    for key in ("trend", "avg7", "avg30", "avg1", "avg", "low"):
        value = _number(row.get(key))
        if value is not None and value > 0:
            return value, key
    return None, None


# The Cardmarket JSON contains far more products than this collection.  We parse
# the file once but keep only Product IDs that belong to matched inventory rows.
def read_relevant_prices(path: str, wanted_ids: set[int]) -> tuple[str, str, dict[int, dict]]:
    """Parse the guide and retain only products already matched to the inventory."""
    # Load Cardmarket's JSON and verify the structure before trusting its fields.
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or not isinstance(data.get("priceGuides"), list):
        raise ValueError("Invalid Cardmarket price guide structure.")
    created_at = str(data.get("createdAt") or "").strip()
    if not created_at:
        raise ValueError("Price guide has no createdAt timestamp.")
    # Store history by calendar date even though Cardmarket provides a full timestamp.
    snapshot_date = datetime.fromisoformat(created_at.replace("Z", "+00:00")).date().isoformat()

    # Only retain products that are actually present in this inventory.
    found: dict[int, dict] = {}
    for row in data["priceGuides"]:
        try:
            pid = int(row.get("idProduct"))
        except (TypeError, ValueError):
            continue
        if pid not in wanted_ids:
            continue
        benchmark, source = _benchmark(row)
        found[pid] = {
            "product_id": pid,
            "snapshot_date": snapshot_date,
            "created_at": created_at,
            "avg": _number(row.get("avg")),
            "low": _number(row.get("low")),
            "trend": _number(row.get("trend")),
            "avg1": _number(row.get("avg1")),
            "avg7": _number(row.get("avg7")),
            "avg30": _number(row.get("avg30")),
            "avg_foil": _number(row.get("avg-foil")),
            "low_foil": _number(row.get("low-foil")),
            "trend_foil": _number(row.get("trend-foil")),
            "avg1_foil": _number(row.get("avg1-foil")),
            "avg7_foil": _number(row.get("avg7-foil")),
            "avg30_foil": _number(row.get("avg30-foil")),
            "benchmark": benchmark,
            "benchmark_source": source,
        }
    return snapshot_date, created_at, found


# Persist one day's relevant product prices and calculate the portfolio total.
# The database key is (product_id, snapshot_date), so running the updater twice
# on the same date updates that date instead of creating duplicate history rows.
def store_snapshot(
    con: sqlite3.Connection,
    inventory: list[dict],
    snapshot_date: str,
    created_at: str,
    prices: dict[int, dict],
) -> dict:
    # Upsert one row per (Product ID, date). Running twice on the same day updates
    # that date instead of creating duplicate card-history points.
    price_sql = """
        INSERT INTO price_history (
            product_id, snapshot_date, created_at, avg, low, trend, avg1, avg7, avg30,
            avg_foil, low_foil, trend_foil, avg1_foil, avg7_foil, avg30_foil,
            benchmark, benchmark_source
        ) VALUES (
            :product_id, :snapshot_date, :created_at, :avg, :low, :trend, :avg1, :avg7, :avg30,
            :avg_foil, :low_foil, :trend_foil, :avg1_foil, :avg7_foil, :avg30_foil,
            :benchmark, :benchmark_source
        )
        ON CONFLICT(product_id, snapshot_date) DO UPDATE SET
            created_at=excluded.created_at, avg=excluded.avg, low=excluded.low,
            trend=excluded.trend, avg1=excluded.avg1, avg7=excluded.avg7,
            avg30=excluded.avg30, avg_foil=excluded.avg_foil, low_foil=excluded.low_foil,
            trend_foil=excluded.trend_foil, avg1_foil=excluded.avg1_foil,
            avg7_foil=excluded.avg7_foil, avg30_foil=excluded.avg30_foil,
            benchmark=excluded.benchmark, benchmark_source=excluded.benchmark_source
    """
    total_value = 0.0
    priced_rows = 0
    unpriced_rows = 0
    total_units = 0
    no_price_rows = 0

    # Calculate the portfolio totals from the selected benchmark of every card row.
    for card in inventory:
        total_units += card["quantity"]
        pid = card.get("product_id")
        price = prices.get(pid) if pid is not None else None
        if price and price["benchmark"] is not None:
            priced_rows += 1
            total_value += card["quantity"] * price["benchmark"]
        else:
            unpriced_rows += 1
            if card.get("match_status") == "MATCHED":
                no_price_rows += 1

    # Save card prices and the portfolio summary together in one transaction.
    with con:
        if prices:
            con.executemany(price_sql, prices.values())
        con.execute(
            """
            INSERT INTO portfolio_history (
                snapshot_date, created_at, total_value, priced_rows, unpriced_rows, total_units
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(snapshot_date) DO UPDATE SET
                created_at=excluded.created_at,
                total_value=excluded.total_value,
                priced_rows=excluded.priced_rows,
                unpriced_rows=excluded.unpriced_rows,
                total_units=excluded.total_units
            """,
            (snapshot_date, created_at, total_value, priced_rows, unpriced_rows, total_units),
        )

    return {
        "snapshot_date": snapshot_date,
        "created_at": created_at,
        "total_value": total_value,
        "priced_rows": priced_rows,
        "unpriced_rows": unpriced_rows,
        "no_price_rows": no_price_rows,
        "total_units": total_units,
        "matched_products_in_guide": len(prices),
    }


# -----------------------------------------------------------------------------
# CENTRAL UPDATE PIPELINE
# -----------------------------------------------------------------------------
# This is the function called by update.py.  In order, it:
#   1) reads Excel,
#   2) resolves/caches Cardmarket Product IDs through matcher.py,
#   3) mirrors the matched inventory into SQLite,
#   4) downloads today's Cardmarket price guide (unless a local guide is supplied),
#   5) extracts only the prices we need,
#   6) stores the product-price snapshot and total portfolio value.
#
# The temporary downloaded price-guide JSON is deleted afterwards; the useful
# history remains in SQLite.
def update(
    price_guide: str | None = None,
    inventory_file: str | None = None,
    rematch: bool = False,
) -> dict:
    """Match inventory, process one Cardmarket snapshot, and return a summary."""
    from matcher import match_inventory

    # Step 1: read the user's human-readable inventory.
    raw_inventory = load_inventory(inventory_file)

    # Step 2: open/create the local SQLite database.
    con = connect()
    init_db(con)
    try:
        # Step 3: attach automatic/manual Cardmarket matches, reusing the cache
        # whenever the card-identifying fields have not changed.
        inventory, match_summary = match_inventory(raw_inventory, con, rematch=rematch)

        # Keep the current enriched inventory in SQLite for the GUI.
        sync_inventory(con, inventory)

        # Build the smallest possible set of Product IDs needed from the huge price guide.
        wanted_ids = {
            int(row["product_id"])
            for row in inventory
            if row.get("product_id") is not None and row.get("match_status") == "MATCHED"
        }

        # Step 4: normal daily operation downloads the current Cardmarket guide.
        # --price-guide is mainly useful for testing/offline reproducibility.
        downloaded = price_guide is None
        guide_path = os.path.abspath(price_guide) if price_guide else download_price_guide()
        try:
            # Steps 5-6: extract our products only, then save this dated snapshot.
            snapshot_date, created_at, prices = read_relevant_prices(guide_path, wanted_ids)
            summary = store_snapshot(con, inventory, snapshot_date, created_at, prices)
        finally:
            if downloaded:
                # The raw guide is temporary; the useful subset is already in SQLite.
                if os.path.exists(guide_path):
                    os.remove(guide_path)

        summary.update(match_summary)
        summary["inventory_file"] = os.path.abspath(inventory_file or inventory_path())
        return summary
    finally:
        con.close()


# -----------------------------------------------------------------------------
# Read-only query helpers used by the GUI
# -----------------------------------------------------------------------------
# These functions do not access the network.  app.py only displays information
# already stored by the most recent updater run.
def latest_portfolio(con: sqlite3.Connection | None = None):
    owns = con is None
    con = con or connect()
    init_db(con)
    row = con.execute(
        "SELECT * FROM portfolio_history ORDER BY snapshot_date DESC LIMIT 1"
    ).fetchone()
    if owns:
        con.close()
    return row


def portfolio_history(con: sqlite3.Connection | None = None):
    owns = con is None
    con = con or connect()
    init_db(con)
    rows = con.execute("SELECT * FROM portfolio_history ORDER BY snapshot_date").fetchall()
    if owns:
        con.close()
    return rows


# Join each current inventory row to the newest stored price for its Product ID.
# This is the main query behind the All cards and Issues GUI tables.
def current_cards(search: str = "", con: sqlite3.Connection | None = None):
    """Return inventory joined to the newest stored price for each product."""
    owns = con is None
    con = con or connect()
    init_db(con)
    query = """
        WITH latest AS (
            SELECT ph.*
            FROM price_history ph
            JOIN (
                SELECT product_id, MAX(snapshot_date) AS snapshot_date
                FROM price_history GROUP BY product_id
            ) x USING (product_id, snapshot_date)
        )
        SELECT i.*, l.snapshot_date, l.trend, l.avg7, l.benchmark, l.benchmark_source,
               CASE WHEN l.benchmark IS NULL THEN NULL ELSE i.quantity * l.benchmark END AS total_value
        FROM inventory i
        LEFT JOIN latest l ON l.product_id = i.product_id
    """
    params: list[str] = []
    if search.strip():
        query += """
            WHERE lower(i.english_name) LIKE ?
               OR lower(i.card_code) LIKE ?
               OR lower(i.rarity) LIKE ?
               OR lower(i.location) LIKE ?
               OR lower(i.match_status) LIKE ?
        """
        term = f"%{search.strip().lower()}%"
        params = [term, term, term, term, term]
    query += " ORDER BY COALESCE(total_value, -1) DESC, lower(i.english_name)"
    rows = con.execute(query, params).fetchall()
    if owns:
        con.close()
    return rows


# Return every stored daily price record for the Product ID currently associated
# with one TRICARD row.  This powers the double-click Card history window.
def card_price_history(card_id: str, con: sqlite3.Connection | None = None):
    """Return every stored Cardmarket price snapshot for one inventory row.

    Price history is keyed by the card's currently resolved Cardmarket product.
    Unmatched cards therefore return an empty history until they are resolved.
    """
    owns = con is None
    con = con or connect()
    init_db(con)
    rows = con.execute(
        """
        SELECT ph.*
        FROM inventory i
        JOIN price_history ph ON ph.product_id = i.product_id
        WHERE i.card_id = ?
        ORDER BY ph.snapshot_date
        """,
        (card_id,),
    ).fetchall()
    if owns:
        con.close()
    return rows
